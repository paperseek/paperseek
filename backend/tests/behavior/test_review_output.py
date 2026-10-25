from __future__ import annotations

import json
import unittest
from datetime import datetime, timezone
from io import BytesIO

from docx import Document
from fixtures.review_support import ReviewCase

from paperseek_app.execution import CancellationScope
from paperseek_app.review.citations import author_year_labels
from paperseek_app.review.exports import export_review
from paperseek_app.review.models import ReviewContent, ReviewDetail
from paperseek_app.review.rendering import render_review
from paperseek_core.json_repair import close_brackets
from paperseek_core.models import FinalPaperResult
from paperseek_core.review.evidence import snapshot, validate_cards, validate_mentions
from paperseek_core.review.models import (
    EvidenceBatch,
    ReviewOutline,
    SectionBody,
    WrittenSection,
)
from paperseek_core.review.structured import (
    ReviewOutputError,
    _clean_json,
    structured_call,
)


class ReviewOutputTests(ReviewCase):
    async def test_schema_is_supplied_first_and_a_complete_json_fence_is_accepted(self):
        class Fenced:
            calls = 0

            async def chat_json(inner, messages, **kwargs):
                inner.calls += 1
                self.assertIn('"maxItems":40', messages[0]["content"])
                self.assertIn('"additionalProperties":false', messages[0]["content"])
                return '```json\n{"paragraphs":[{"text":"Finding","evidence_ids":["e_known"]}]}\n```'

        llm = Fenced()
        body = await structured_call(llm, CancellationScope(), "section", "Return paragraphs.", {}, SectionBody)
        self.assertEqual(body.paragraphs[0].text, "Finding")
        self.assertEqual(llm.calls, 1)

    async def test_invalid_json_is_not_salvaged_and_failure_retains_the_task_and_field(self):
        class Invalid:
            calls = 0

            async def chat_json(inner, messages, **kwargs):
                inner.calls += 1
                if inner.calls == 1:
                    return 'Explanation: {"paragraphs": []}'
                return '{"paragraphs":[{"text":["private-input-value"],"evidence_ids":["e_known"]}]}'

        llm = Invalid()
        with self.assertRaises(ReviewOutputError) as caught:
            await structured_call(llm, CancellationScope(), "section_edit", "Return paragraphs.", {}, SectionBody)
        self.assertEqual(llm.calls, 2)
        self.assertEqual(caught.exception.task, "section_edit")
        self.assertIn("paragraphs.0.text", caught.exception.reason)
        self.assertNotIn("private-input-value", caught.exception.reason)

    async def test_a_reply_missing_its_final_brackets_is_closed_without_another_model_call(self):
        class Unclosed:
            calls = 0

            async def chat_json(inner, messages, **kwargs):
                inner.calls += 1
                # Intern-S2 sometimes ends a complete reply without its last "}".
                return '{"paragraphs":[{"text":"A brace } inside the text","evidence_ids":["e_known"]}]'

        llm = Unclosed()
        body = await structured_call(llm, CancellationScope(), "section", "", {}, SectionBody)
        self.assertEqual(body.paragraphs[0].text, "A brace } inside the text")
        self.assertEqual(llm.calls, 1)
        # A reply cut off inside a string is not completed by guesswork.
        self.assertEqual(close_brackets('{"paragraphs":[{"text":"cut'), '{"paragraphs":[{"text":"cut')

    async def test_a_fenced_reply_missing_its_final_brackets_keeps_every_paragraph(self):
        class Fenced:
            async def chat_json(inner, messages, **kwargs):
                return ('```json' + chr(10) + '{"paragraphs":[{"text":"First","evidence_ids":["e_1"]},'
                        '{"text":"Second","evidence_ids":["e_2"]' + chr(10) + '```')

        body = await structured_call(Fenced(), CancellationScope(), "section", "", {}, SectionBody)
        self.assertEqual([paragraph.text for paragraph in body.paragraphs], ["First", "Second"])

    async def test_a_reply_cut_off_inside_a_later_paragraph_is_repaired_not_shortened(self):
        complete = '{"paragraphs":[{"text":"First","evidence_ids":["e_1"]},{"text":"Second","evidence_ids":["e_2"]}]}'
        calls = []

        class CutOff:
            async def chat_json(inner, messages, **kwargs):
                calls.append(messages[0]["content"])
                if messages[0]["content"].startswith("You are a strict JSON formatter"):
                    return complete
                return '{"paragraphs":[{"text":"First","evidence_ids":["e_1"]},{"text":"Second is cut mid'

        body = await structured_call(CutOff(), CancellationScope(), "section", "", {}, SectionBody)
        self.assertEqual([paragraph.text for paragraph in body.paragraphs], ["First", "Second"])
        self.assertEqual(len(calls), 2)
        # Explanatory text after a complete object is still removed without another call.
        self.assertEqual(_clean_json(complete + " Hope this helps."), complete)

    async def test_broken_json_is_reformatted_without_regenerating_content(self):
        calls = []
        valid = json.dumps({"paragraphs": [{"text": 'He said "yes".', "evidence_ids": ["e_1"]}]})

        class Model:
            async def chat_json(self, messages, **_):
                calls.append(messages[0]["content"])
                if messages[0]["content"].startswith("You are a strict JSON formatter"):
                    return valid
                # Unescaped inner quotes: the content is fine, the JSON encoding is not.
                return '{"paragraphs":[{"text":"He said "yes".","evidence_ids":["e_1"]}]}'

        body = await structured_call(Model(), CancellationScope(), "section", "", {}, SectionBody)
        self.assertEqual(body.paragraphs[0].text, 'He said "yes".')
        self.assertEqual(len(calls), 2)
        self.assertTrue(calls[1].startswith("You are a strict JSON formatter"))


class ReviewCitationsTest(unittest.TestCase):
    def setUp(self):
        self.papers = [snapshot(FinalPaperResult(rank=i, uid=str(i), title=title,
            authors=authors, year=year, abstract="Original study finding."), "source", i)
            for i, title, authors, year in [
                (1, "B study", ["Jane Smith"], 2020), (2, "A study", ["Jane Smith"], 2020),
                (3, "C study", [], None), (4, "D study", ["王明", "李华"], 2024)]]
        self.cards = validate_cards(EvidenceBatch(cards=[{"paper_id": p.id, "evidence": [
            {"category": "finding", "text": "Finding", "quote": p.abstract}]} for p in self.papers]), self.papers, "test")
        self.keys = [card.evidence[0].id for card in self.cards]
        body = SectionBody(paragraphs=[{"text": "p_002 and p_001 agree (p_002、p_001).", "evidence_ids": self.keys[:2]},
                                       {"text": "A further finding.", "evidence_ids": self.keys[2:]}])
        content = ReviewContent(topic="Compare findings", language="en", length="short", papers=self.papers,
            corpus_hash="test", cards=self.cards, outline=ReviewOutline(sections=[
                {"id": "s1", "title": "Findings", "goal": "Compare", "paper_ids": [p.id for p in self.papers]}]),
            sections=[WrittenSection(section_id="s1", revision=1, body=body, author="model", signature="test",
                                     validation="supported")])
        self.review = ReviewDetail(id="review", source_run_id="source", version=1, content=content,
                                   created_at=datetime.now(timezone.utc), updated_at=datetime.now(timezone.utc))

    def test_numeric_order_follows_inline_mentions_and_does_not_repeat_citations_at_end(self):
        document = render_review(self.review)
        paragraph = document.sections[0].paragraphs[0]
        self.assertEqual([r.paper_id for r in document.references], ["p_002", "p_001", "p_003", "p_004"])
        self.assertEqual("".join(s.citation.label if s.citation else s.text for s in paragraph.spans),
                         "[1] and [2] agree [1]; [2].")
        self.assertEqual([s.citation.paper_id for s in paragraph.spans if s.citation],
                         ["p_002", "p_001", "p_002", "p_001"])

    def test_unplaced_citations_are_listed_apart_and_evidence_ids_never_reach_the_reader(self):
        section = self.review.content.sections[0]
        section.body.paragraphs[0].text = "p_002 and p_001 agree (p_002、p_001) (e_0123456789abcdef01234567)."
        document = render_review(self.review)
        first, second = document.sections[0].paragraphs
        self.assertNotIn("e_0123", first.text)
        self.assertEqual([citation.paper_id for citation in second.unplaced], ["p_003", "p_004"])
        self.assertFalse([span for span in second.spans if span.citation])
        self.assertEqual(document.sections[0].validation, "needs_revision")
        markdown = export_review(self.review, "md").content.decode()
        self.assertIn("Also cited (not placed beside a claim): [3](#ref-3); [4](#ref-4)", markdown)
        docx = "\n".join(p.text for p in Document(BytesIO(export_review(self.review, "docx").content)).paragraphs)
        self.assertIn("Also cited (not placed beside a claim): [3]; [4]", docx)

    def test_harvard_uses_metadata_disambiguates_and_sorts_references(self):
        document = render_review(self.review, "harvard")
        paragraph = document.sections[0].paragraphs[0]
        labels = [s.citation.label for s in paragraph.spans if s.citation]
        self.assertEqual(labels, ["(Smith, 2020a)", "(Smith, 2020b)"] * 2)
        self.assertIn("(2020a)", next(r.text for r in document.references if r.paper_id == "p_002"))
        self.assertEqual(author_year_labels(self.papers)["p_003"], ("“C study”", "n.d."))
        self.assertEqual(author_year_labels(self.papers)["p_004"], ("王明 and 李华", "2024"))
        many = self.papers[0].model_copy(update={"authors": ["Jane Smith", "Jones, Lee", "Anna de la Cruz"]})
        self.assertEqual(author_year_labels([many])[many.id], ("Smith et al.", "2020"))

    def test_unbound_inline_paper_is_flagged_instead_of_creating_a_false_citation(self):
        paragraph = self.review.content.sections[0].body.paragraphs[0]
        paragraph.text = "p_004 found something."
        with self.assertRaisesRegex(ValueError, "without evidence"):
            validate_mentions(self.review.content.sections[0].body, self.cards)
        document = render_review(self.review)
        rendered = document.sections[0].paragraphs[0]
        self.assertTrue(document.draft)
        self.assertEqual(document.sections[0].validation, "needs_revision")
        self.assertEqual(len(rendered.citation_warnings), 1)
        self.assertNotIn("p_004", "".join(s.text for s in rendered.spans))
        self.assertNotIn("p_004", [s.citation.paper_id for s in rendered.spans if s.citation])
        self.assertIn("citation needs checking", export_review(self.review, "md").content.decode())

    def test_cited_papers_must_be_named_beside_their_claims_not_stacked_at_the_end(self):
        body = self.review.content.sections[0].body
        body.paragraphs[1].text = "A further finding."
        with self.assertRaisesRegex(ValueError, "without naming"):
            validate_mentions(body, self.cards)
        body.paragraphs[1].text = "p_003 and p_004 report further findings."
        body.paragraphs[0].evidence_ids = self.keys[:3]
        body.paragraphs[0].text = "Recall improves. Precision holds. Both agree (p_001, p_002, p_003)."
        with self.assertRaisesRegex(ValueError, "groups all its citations at the end"):
            validate_mentions(body, self.cards)
        body.paragraphs[0].text = "召回率提高（p_001），精确率保持稳定（p_002），两者一致（p_003）。"
        validate_mentions(body, self.cards)

    def test_chinese_paragraphs_reject_english_sentences_but_allow_a_first_mention_gloss(self):
        body = self.review.content.sections[0].body
        body.paragraphs[1].text = "研究报告了相关发现（p_003），大语言模型（Large Language Model, LLM）也有类似结果（p_004）。"
        validate_mentions(body, self.cards, language="zh")
        body.paragraphs[1].text = ("研究（p_003）指出：Extended Reality supports situated human control in "
                                   "dynamic shared environments for robots（p_004）。")
        with self.assertRaisesRegex(ValueError, "contains English"):
            validate_mentions(body, self.cards, language="zh")
        validate_mentions(body, self.cards, language="en")
        validate_mentions(body, self.cards, language="zh", exempt={body.paragraphs[1].text})

    def test_chinese_reviews_are_prompted_in_chinese(self):
        from paperseek_core.review.prompts import (
            SYSTEM,
            SYSTEM_ZH,
            _say,
        )
        self.assertIn("大语言模型（Large Language Model, LLM）", SYSTEM_ZH)
        self.assertEqual(_say("zh", SYSTEM, SYSTEM_ZH), SYSTEM_ZH)
        self.assertEqual(_say("en", SYSTEM, SYSTEM_ZH), SYSTEM)

    def test_exports_share_the_view_labels_without_modifying_the_saved_draft(self):
        before = self.review.model_dump_json()
        for style, label in [("numeric", "[1]"), ("harvard", "(Smith, 2020a)")]:
            markdown = export_review(self.review, "md", style).content.decode()
            docx = Document(BytesIO(export_review(self.review, "docx", style).content))
            self.assertIn(label, markdown)
            self.assertIn(label, "\n".join(p.text for p in docx.paragraphs))
            self.assertNotIn("p_002", markdown)
            self.assertNotIn("p_002", "\n".join(p.text for p in docx.paragraphs))
        self.assertIn("paperseek_review_p_002", export_review(self.review, "bibtex", "harvard").content.decode())
        self.assertEqual(self.review.model_dump_json(), before)
