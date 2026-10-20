from __future__ import annotations

import hashlib
from html import escape
from io import BytesIO
import re
from urllib.parse import urlsplit

from docx import Document

from paperseek_app.models import ArtifactContent
from paperseek_app.review.models import ReviewDetail, ReviewError
from paperseek_app.review.rendering import render_review
from paperseek_app.review.citations import CitationStyle


def _text(value: str) -> str:
    return re.sub(r"([\\`*_\[\]#])", r"\\\1", escape(value, quote=False))


def _bib(value: str) -> str:
    return "".join({"\\": r"\textbackslash{}", "{": r"\{", "}": r"\}", "%": r"\%", "&": r"\&",
                    "_": r"\_", "#": r"\#", "$": r"\$", "~": r"\~{}", "^": r"\^{}"}.get(c, c) for c in value)


def export_review(review: ReviewDetail, format: str, citation_style: CitationStyle = "numeric") -> ArtifactContent:
    rendered = render_review(review, citation_style)
    if not rendered.sections:
        raise ReviewError("review_no_document", status=422)
    zh = rendered.language == "zh"
    label = "摘要证据综述" if zh else "Abstract-based literature review"
    if rendered.draft:
        label += " · 草稿，含待检查或待更新内容" if zh else " · Draft with pending checks or updates"
    heading = "参考文献" if zh else "References"
    unplaced_label = "本段另引（未对应到具体论断）：" if zh else "Also cited (not placed beside a claim):"
    if format == "md":
        lines = ["# " + _text(rendered.title), "", label, ""]
        for section in rendered.sections:
            lines.extend(["#" * (section.level + 1) + " " + _text(section.title), ""])
            for paragraph in section.paragraphs:
                parts = [f"[{c.number if citation_style == 'numeric' else _text(c.label)}](#ref-{c.number})"
                         if (c := span.citation) else _text(span.text) for span in paragraph.spans]
                lines.extend(["".join(parts), ""])
                if paragraph.unplaced:
                    lines.extend(["*" + unplaced_label + " " + "; ".join(
                        f"[{c.number if citation_style == 'numeric' else _text(c.label)}](#ref-{c.number})"
                        for c in paragraph.unplaced) + "*", ""])
                lines.extend("> " + _text(warning) + "\n" for warning in paragraph.citation_warnings)
        lines.extend(["## " + heading, ""])
        for reference in rendered.references:
            link = reference.url if urlsplit(reference.url).scheme in {"https", "http"} else ""
            lines.extend([f'<a id="ref-{reference.number}"></a>',
                          (f"[{reference.number}] " if citation_style == "numeric" else "")
                          + _text(reference.text) + (" " + _text(link) if link else ""), ""])
        content, media = "\n".join(lines).encode(), "text/markdown; charset=utf-8"
    elif format == "docx":
        document = Document()
        document.add_heading(rendered.title, 0)
        document.add_paragraph(label)
        for section in rendered.sections:
            document.add_heading(section.title, section.level)
            for paragraph in section.paragraphs:
                document.add_paragraph("".join(span.citation.label if span.citation else span.text for span in paragraph.spans))
                if paragraph.unplaced:
                    document.add_paragraph(unplaced_label + " " + "; ".join(c.label for c in paragraph.unplaced))
                for warning in paragraph.citation_warnings:
                    document.add_paragraph(warning)
        document.add_heading(heading, 1)
        for reference in rendered.references:
            document.add_paragraph((f"[{reference.number}] " if citation_style == "numeric" else "")
                                   + reference.text + (f" DOI: {reference.doi}" if reference.doi else ""))
        output = BytesIO()
        document.save(output)
        content, media = output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif format == "bibtex":
        papers = {paper.id: paper for paper in review.content.papers}
        entries = []
        for reference in rendered.references:
            paper = papers[reference.paper_id]
            fields = {"title": paper.title, "author": " and ".join(paper.authors), "year": str(paper.year or ""),
                      "howpublished": paper.venue, "doi": paper.doi, "url": paper.url}
            entries.append(f"@misc{{paperseek_{review.id[:8]}_{paper.id},\n" + ",\n".join(
                f"  {key} = {{{_bib(value)}}}" for key, value in fields.items() if value) + "\n}")
        content, media = ("\n\n".join(entries) + "\n").encode(), "application/x-bibtex"
    else:
        raise ReviewError("invalid_export_format", status=422)
    extension = "bib" if format == "bibtex" else format
    return ArtifactContent(id=review.id, filename=f"review-{review.id}-v{review.version}.{extension}",
                           media_type=media, content=content, size=len(content),
                           sha256=hashlib.sha256(content).hexdigest())
